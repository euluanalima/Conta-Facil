# -*- coding: utf-8 -*-
"""
CONTA FÁCIL - NÚCLEO DE REGRAS
Funções reaproveitadas da aplicação desktop para:
  1) leitura de PDF consolidado da EMBASA;
  2) consolidação no padrão histórico de controle da UFRB;
  3) geração de Excel (.xlsx);
  4) gráficos analíticos integrados e exportação PDF;
  5) geração de Relatório Mensal de Ateste em Word (.docx).

Dependências:
    pip install PyQt5 pandas matplotlib pypdf python-docx xlsxwriter

Execução:
    python Conta_Facil.py

Observação sobre precisão monetária:
Todos os valores financeiros são convertidos para CENTAVOS (inteiros) no momento
da extração. As somas/agrupamentos do pandas trabalham com inteiros, evitando erros
de ponto flutuante. Conversões para reais ocorrem apenas na apresentação/exportação.
"""

# ============================================================================
# GUIA DE LEITURA DO CÓDIGO
# ============================================================================
# Este arquivo foi mantido como aplicação monolítica (um único .py) para facilitar
# distribuição e manutenção em ambiente administrativo. A lógica está organizada
# em cinco camadas informais:
#
# 1) CONFIGURAÇÃO E BASE HISTÓRICA
#    Constantes, nomes de centros, meses e dados históricos incorporados.
# 2) EXTRAÇÃO E TRATAMENTO DE DADOS
#    Leitura do PDF da EMBASA, conversão monetária e normalização do histórico.
# 3) CONSOLIDAÇÃO E ANÁLISE
#    Associação por matrícula, totais, percentuais e variações mensais.
# 4) EXPORTAÇÕES
#    Geração de Excel, gráficos/PDF e relatório Word de ateste.
# 5) INTERFACE GRÁFICA
#    Classes Qt responsáveis por tabela, gráficos, filtros e ações do usuário.
#
# FLUXO PRINCIPAL DE DADOS:
# PDF EMBASA -> extract_pdf() -> DataFrame de faturas -> build_control()
# -> DataFrame de controle + resumo por centro -> interface / Excel / gráficos / Word
#
# REGRA IMPORTANTE DE PRECISÃO:
# valores monetários são mantidos em CENTAVOS (inteiros) durante cálculos. A
# conversão para reais ocorre apenas na apresentação e nas exportações. Isso evita
# erros de ponto flutuante em somas financeiras.
# ============================================================================

from __future__ import annotations

import sys
import os
import re
import json
from pathlib import Path
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import datetime
from typing import Optional, Tuple, List, Dict

# Dados tabulares e leitura do PDF consolidado.
import pandas as pd
from pypdf import PdfReader

# Camada de interface gráfica (PyQt5).

# Visualização e exportação de gráficos.
from matplotlib.figure import Figure
from matplotlib.backends.backend_pdf import PdfPages

# Geração e formatação do relatório de ateste em Word.
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# Constantes de domínio usadas em toda a aplicação.
APP_TITLE = "Conta Fácil"
CENTROS_ORDEM = ["CDA", "CFP", "CCS", "CAHL", "CECULT", "CETENS", "NÃO MAPEADO"]
MESES_PT = {1:"Jan", 2:"Fev", 3:"Mar", 4:"Abr", 5:"Mai", 6:"Jun", 7:"Jul", 8:"Ago", 9:"Set", 10:"Out", 11:"Nov", 12:"Dez"}
MESES_EXTENSO = {1:"janeiro", 2:"fevereiro", 3:"março", 4:"abril", 5:"maio", 6:"junho", 7:"julho", 8:"agosto", 9:"setembro", 10:"outubro", 11:"novembro", 12:"dezembro"}

# Base histórica reconstruída a partir da planilha presente no Relatório Mensal de
# Ateste de junho/2026 fornecido pela UFRB. Valores armazenados em centavos.
BASE_HISTORICA = [
    {
        "ordem": 1,
        "matricula": "57162735",
        "local": "AV DR LUIS SANDES, 190, AP 08, SANTA RTA",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4270,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 2,
        "matricula": "180552090",
        "local": "GINÁSIO DE ESPORTES",
        "centro": "CFP",
        "Mar/2026": 146590,
        "Abr/2026": 98510,
        "Mai/2026": 58540,
        "Jun/2026": 138430
    },
    {
        "ordem": 3,
        "matricula": "57162794",
        "local": "AV DR LUIS SANDES, 190, AP 14, S,T,",
        "centro": "CFP",
        "Mar/2026": 4180,
        "Abr/2026": 4400,
        "Mai/2026": 4480,
        "Jun/2026": 4320
    },
    {
        "ordem": 4,
        "matricula": "57162786",
        "local": "AV DR LUIS SANDES, 190, AP 13, S, T,",
        "centro": "CFP",
        "Mar/2026": 4180,
        "Abr/2026": 4280,
        "Mai/2026": 10580,
        "Jun/2026": 5960
    },
    {
        "ordem": 5,
        "matricula": "57162778",
        "local": "AV DR LUISA SANDES, 190, AP 12, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4270,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 6,
        "matricula": "57162760",
        "local": "AV DR LUIS SANDES, 190, AP 11, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 7,
        "matricula": "57162751",
        "local": "AV DR LUIS SANDES,190, AP 10, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 4320,
        "Jun/2026": 4150
    },
    {
        "ordem": 8,
        "matricula": "57162743",
        "local": "AV DR LUISA SANDES, 190, AP 09, S, T,",
        "centro": "CFP",
        "Mar/2026": 4210,
        "Abr/2026": 4370,
        "Mai/2026": 4350,
        "Jun/2026": 4150
    },
    {
        "ordem": 9,
        "matricula": "57162727",
        "local": "AV DR LUIS SANDES, 190, AP 07, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4420,
        "Mai/2026": 4480,
        "Jun/2026": 4150
    },
    {
        "ordem": 10,
        "matricula": "57162719",
        "local": "AV DR LUIS SANDES, 190, AP 06, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4270,
        "Mai/2026": 4650,
        "Jun/2026": 4150
    },
    {
        "ordem": 11,
        "matricula": "57162701",
        "local": "AV DR LUIS SANDES, 190, AP 05, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 7120,
        "Jun/2026": 4640
    },
    {
        "ordem": 12,
        "matricula": "57162697",
        "local": "AV DR LUIS SANDES, 190, AP 04, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 13,
        "matricula": "57162689",
        "local": "AV DR LUIS SANDES, 190, AP 03, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 14,
        "matricula": "57162671",
        "local": "AV DR LUIS SANDES, 190, AP 02",
        "centro": "CFP",
        "Mar/2026": 4180,
        "Abr/2026": 4260,
        "Mai/2026": 8270,
        "Jun/2026": 8270
    },
    {
        "ordem": 15,
        "matricula": "181973073",
        "local": "PÇ IRACI SILVA, 55, CASA DO DUCA",
        "centro": "CFP",
        "Mar/2026": 12110,
        "Abr/2026": 12370,
        "Mai/2026": 12070,
        "Jun/2026": 12510
    },
    {
        "ordem": 16,
        "matricula": "57094314",
        "local": "AV, NESTOR DE MELO PITA, SEDE DO CFP",
        "centro": "CFP",
        "Mar/2026": 511582,
        "Abr/2026": 389612,
        "Mai/2026": 908640,
        "Jun/2026": 886670
    },
    {
        "ordem": 17,
        "matricula": "57049521",
        "local": "AV DR LUIS SANDES, 190, SANTA RITA",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4260,
        "Mai/2026": 4160,
        "Jun/2026": 4150
    },
    {
        "ordem": 18,
        "matricula": "57029997",
        "local": "RUA CEL BENEDITO ALMEIDA, TECELENDO",
        "centro": "CFP",
        "Mar/2026": 11960,
        "Abr/2026": 13540,
        "Mai/2026": 12980,
        "Jun/2026": 12050
    },
    {
        "ordem": 19,
        "matricula": "57162662",
        "local": "AV DR LUIS SANDES, 190, AP 01, S, T,",
        "centro": "CFP",
        "Mar/2026": 4170,
        "Abr/2026": 4270,
        "Mai/2026": 8270,
        "Jun/2026": 4150
    },
    {
        "ordem": 20,
        "matricula": "75416271",
        "local": "RUA JJ SEABRA, 5, PAV, LEITE ALVES",
        "centro": "CAHL",
        "Mar/2026": 80930,
        "Abr/2026": 260450,
        "Mai/2026": 616190,
        "Jun/2026": 409240
    },
    {
        "ordem": 21,
        "matricula": "75386402",
        "local": "RUA 13 DE MAIO, FUND, HANSEN",
        "centro": "CAHL",
        "Mar/2026": 169440,
        "Abr/2026": 120230,
        "Mai/2026": 176950,
        "Jun/2026": 153650
    },
    {
        "ordem": 22,
        "matricula": "75385716",
        "local": "RUA ANA NERY, 9, CENTRO, NUDOC",
        "centro": "CAHL",
        "Mar/2026": 21770,
        "Abr/2026": 22260,
        "Mai/2026": 21740,
        "Jun/2026": 21700
    },
    {
        "ordem": 23,
        "matricula": "75385830",
        "local": "RUA ANA NERY, 25 A, PROCULTURA",
        "centro": "CAHL",
        "Mar/2026": 23420,
        "Abr/2026": 23100,
        "Mai/2026": 24190,
        "Jun/2026": 24180
    },
    {
        "ordem": 24,
        "matricula": "96153970",
        "local": "PRÉDIO DA PÓS-GRADUAÇÃO DO CAHL",
        "centro": "CAHL",
        "Mar/2026": 21530,
        "Abr/2026": 100420,
        "Mai/2026": 86500,
        "Jun/2026": 105120
    },
    {
        "ordem": 25,
        "matricula": "96151811",
        "local": "PAVILHÃO 2 DE JULHO",
        "centro": "CAHL",
        "Mar/2026": 21530,
        "Abr/2026": 23310,
        "Mai/2026": 72260,
        "Jun/2026": 22520
    },
    {
        "ordem": 26,
        "matricula": "96176288",
        "local": "RESID, ESTUDANTIL ADEMIR FERNADO",
        "centro": "CAHL",
        "Mar/2026": 53150,
        "Abr/2026": 43410,
        "Mai/2026": 71770,
        "Jun/2026": 71920
    },
    {
        "ordem": 27,
        "matricula": "54492866",
        "local": "RESIDÊNCIA ESTUDANTIL DA CIDADE",
        "centro": "CDA",
        "Mar/2026": 7420,
        "Abr/2026": 14400,
        "Mai/2026": 53800,
        "Jun/2026": 93940
    },
    {
        "ordem": 28,
        "matricula": "54469511",
        "local": "GUARITA TABELA",
        "centro": "CDA",
        "Mar/2026": 11960,
        "Abr/2026": 12500,
        "Mai/2026": 12070,
        "Jun/2026": 12050
    },
    {
        "ordem": 29,
        "matricula": "54465494",
        "local": "HIDRÔMETRO DO CETEP",
        "centro": "CDA",
        "Mar/2026": 4611570,
        "Abr/2026": 3365380,
        "Mai/2026": 3249600,
        "Jun/2026": 3456880
    },
    {
        "ordem": 30,
        "matricula": "54465478",
        "local": "GUARITA CETEP",
        "centro": "CDA",
        "Mar/2026": 11960,
        "Abr/2026": 12500,
        "Mai/2026": 12070,
        "Jun/2026": 12050
    },
    {
        "ordem": 31,
        "matricula": "180901141",
        "local": "CAPRINOS",
        "centro": "CDA",
        "Mar/2026": 71520,
        "Abr/2026": 68500,
        "Mai/2026": 40030,
        "Jun/2026": 135200
    },
    {
        "ordem": 32,
        "matricula": "180901168",
        "local": "AVES",
        "centro": "CDA",
        "Mar/2026": 37410,
        "Abr/2026": 14130,
        "Mai/2026": 21660,
        "Jun/2026": 22050
    },
    {
        "ordem": 33,
        "matricula": "54676304",
        "local": "PAVILHÃO DE AULAS II (UFRB)",
        "centro": "CDA",
        "Mar/2026": 23930,
        "Abr/2026": 152040,
        "Mai/2026": 244370,
        "Jun/2026": 349040
    },
    {
        "ordem": 34,
        "matricula": "54690960",
        "local": "LAB BIOLOGIA (UFRB)",
        "centro": "CDA",
        "Mar/2026": 134220,
        "Abr/2026": 218700,
        "Mai/2026": 228080,
        "Jun/2026": 88290
    },
    {
        "ordem": 35,
        "matricula": "54706360",
        "local": "NEAS",
        "centro": "CDA",
        "Mar/2026": 95700,
        "Abr/2026": 102320,
        "Mai/2026": 96520,
        "Jun/2026": 96430
    },
    {
        "ordem": 36,
        "matricula": "87553414",
        "local": "ZOOTECNIA – SAPUCAIA (UFRB)",
        "centro": "CDA",
        "Mar/2026": 114900,
        "Abr/2026": 124570,
        "Mai/2026": 90130,
        "Jun/2026": 153230
    },
    {
        "ordem": 37,
        "matricula": "54672937",
        "local": "PAVILÃO DE AULAS I (UFRB)",
        "centro": "CDA",
        "Mar/2026": 207487,
        "Abr/2026": 225780,
        "Mai/2026": 286870,
        "Jun/2026": 378070
    },
    {
        "ordem": 38,
        "matricula": "97437891",
        "local": "AV CENTENÁRIO, 697, SIM",
        "centro": "CETENS",
        "Mar/2026": 12090,
        "Abr/2026": 40620,
        "Mai/2026": 69000,
        "Jun/2026": 34870
    },
    {
        "ordem": 39,
        "matricula": "96841877",
        "local": "RUA VISC RIO BRANCO(DO), 1 A, CENTRO",
        "centro": "CETENS",
        "Mar/2026": 24820,
        "Abr/2026": 24540,
        "Mai/2026": 48640,
        "Jun/2026": 86410
    },
    {
        "ordem": 40,
        "matricula": "51130076",
        "local": "AV CARLOS AMARAL, 1051, CAJUEIRO, GUARITA",
        "centro": "CCS",
        "Mar/2026": 49240,
        "Abr/2026": 133010,
        "Mai/2026": 167320,
        "Jun/2026": 239710
    },
    {
        "ordem": 41,
        "matricula": "50922270",
        "local": "AV CARLOS AMARAL, 897, CAJUEIRO, PSICOLOGIA",
        "centro": "CCS",
        "Mar/2026": 170064,
        "Abr/2026": 152528,
        "Mai/2026": 252250,
        "Jun/2026": 531930
    },
    {
        "ordem": 42,
        "matricula": "51189500",
        "local": "AV CARLOS AMARAL, 905 - SEDE DO CCS",
        "centro": "CCS",
        "Mar/2026": 105630,
        "Abr/2026": 15010,
        "Mai/2026": 11960,
        "Jun/2026": 13460
    },
    {
        "ordem": 43,
        "matricula": "94250006",
        "local": "RUA IMPERADOR, 5, ESCOLA MUNICIPAL",
        "centro": "CECULT",
        "Mar/2026": 22590,
        "Abr/2026": 62120,
        "Mai/2026": 58540,
        "Jun/2026": 58570
    },
    {
        "ordem": 44,
        "matricula": "94217149",
        "local": "AV VIANA BANDEIRA, 119, CENTRO",
        "centro": "CECULT",
        "Mar/2026": 12090,
        "Abr/2026": 12370,
        "Mai/2026": 12070,
        "Jun/2026": 12060
    }
]

# Colunas fixas que não representam meses do histórico.
STATIC_COLS = {"Nº G", "LOCAL", "CEN", "MAT", "Aum / Red", "% Centro", "% Total"}


# ---------------------------------------------------------------------------
# Converte textos monetários de diferentes formatos para centavos inteiros. Centralizar essa conversão evita inconsistência entre PDF, Excel e relatórios.
# ---------------------------------------------------------------------------
def cents_from_text(value: str) -> int:
    """Converte formatos 1.234,56 / 1234.56 / 1234,56 para centavos exatos."""
    if value is None:
        raise ValueError("Valor monetário ausente")
    s = str(value).strip().replace("R$", "").replace(" ", "")
    if not s:
        raise ValueError("Valor monetário vazio")
    # Quando há vírgula, ela é o separador decimal brasileiro.
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        # Pypdf costuma devolver 2265.45. Se houver múltiplos pontos, tratamos
        # todos menos o último como milhar.
        if s.count(".") > 1:
            parts = s.split(".")
            s = "".join(parts[:-1]) + "." + parts[-1]
    try:
        dec = Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f"Valor inválido: {value}") from exc
    return int((dec * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


# ---------------------------------------------------------------------------
# Transforma centavos em Decimal de reais apenas quando é necessário apresentar/exportar o valor.
# ---------------------------------------------------------------------------
def reais_from_cents(cents: int | float | None) -> Decimal:
    if cents is None or pd.isna(cents):
        return Decimal("0.00")
    return (Decimal(int(cents)) / Decimal(100)).quantize(Decimal("0.01"))


# ---------------------------------------------------------------------------
# Formata um valor em centavos no padrão monetário brasileiro usado na interface e nos relatórios.
# ---------------------------------------------------------------------------
def brl(cents: int | float | None) -> str:
    d = reais_from_cents(cents)
    raw = f"{d:,.2f}"
    return "R$ " + raw.replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------------------
# Formata percentuais para exibição humana, preservando o valor numérico original para cálculos.
# ---------------------------------------------------------------------------
def pct(value: float | Decimal | None) -> str:
    if value is None or pd.isna(value):
        return "-"
    return f"{float(value):.1f}%".replace(".", ",")


# ---------------------------------------------------------------------------
# Converte MM/AAAA para o rótulo interno de coluna (ex.: 07/2026 -> Jul/2026).
# ---------------------------------------------------------------------------
def month_key(mm_yyyy: str) -> str:
    m, y = [int(x) for x in mm_yyyy.split("/")]
    return f"{MESES_PT[m]}/{y}"


# ---------------------------------------------------------------------------
# Produz o nome do mês por extenso para títulos e textos do relatório.
# ---------------------------------------------------------------------------
def month_long(mm_yyyy: str) -> str:
    m, y = [int(x) for x in mm_yyyy.split("/")]
    return f"{MESES_EXTENSO[m]} de {y}"




# ---------------------------------------------------------------------------
# Calcula o mês calendário imediatamente anterior; essa regra é usada nas comparações mensais.
# ---------------------------------------------------------------------------
def previous_month_label(mm_yyyy: str) -> str:
    m, y = [int(x) for x in mm_yyyy.split("/")]
    if m == 1:
        m, y = 12, y - 1
    else:
        m -= 1
    return f"{MESES_PT[m]}/{y}"


# ---------------------------------------------------------------------------
# Faz a conversão inversa do rótulo interno de mês para MM/AAAA.
# ---------------------------------------------------------------------------
def month_label_to_mm_yyyy(label: str) -> str:
    abbr, y = label.split("/", 1)
    rev = {v:k for k,v in MESES_PT.items()}
    return f"{rev[abbr]:02d}/{int(y)}"


# ---------------------------------------------------------------------------
# Fornece uma chave cronológica para ordenar as colunas mensais corretamente.
# ---------------------------------------------------------------------------
def month_sort_key(label: str) -> Tuple[int, int]:
    """Ordena rótulos Jan/2026, Fev/2026..."""
    if "/" not in str(label):
        return (9999, 99)
    abbr, year = str(label).split("/", 1)
    rev = {v:k for k,v in MESES_PT.items()}
    return (int(year), rev.get(abbr, 99))


# ---------------------------------------------------------------------------
# Tenta recuperar o nome associado à conta a partir das linhas extraídas da própria página do PDF.
# ---------------------------------------------------------------------------
def extract_account_name(lines: List[str]) -> str:
    try:
        idx_mes = lines.index("Mês/Ano")
        idx_em = lines.index("Data Emissão")
        if idx_em > idx_mes + 1:
            candidates = [x for x in lines[idx_mes+1:idx_em] if not re.fullmatch(r"\d{2}/\d{4}", x)]
            return " ".join(candidates).strip()
    except ValueError:
        pass
    return ""


# ---------------------------------------------------------------------------
# Núcleo de leitura das faturas: percorre cada página, extrai matrícula, mês, datas e valor, e entrega um DataFrame padronizado.
# ---------------------------------------------------------------------------
def extract_pdf(pdf_path: str) -> pd.DataFrame:
    """Extrai uma fatura por página do PDF consolidado EMBASA."""
    reader = PdfReader(pdf_path)
    records: List[Dict] = []

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        lines = [re.sub(r"\s+", " ", x).strip() for x in text.splitlines() if x.strip()]
        if not lines:
            continue

        # Matrícula: primeiro número de 7 a 9 dígitos antes do título DANFE.
        matricula = None
        danfe_idx = next((i for i, x in enumerate(lines) if "DOCUMENTO AUXILIAR" in x.upper()), min(len(lines), 15))
        for x in lines[:danfe_idx]:
            if re.fullmatch(r"\d{7,9}", x):
                matricula = x
                break
        if not matricula:
            m = re.search(r"N[º°]?\s*DA\s*MATR[ÍI]CULA[\s\S]{0,300}?\b(\d{7,9})\b", text, re.I)
            matricula = m.group(1) if m else None

        # Mês/ano: primeira ocorrência MM/AAAA na área de cabeçalho.
        header = "\n".join(lines[:40])
        mm = re.search(r"\b(0[1-9]|1[0-2])/20\d{2}\b", header)
        mes_ano = mm.group(0) if mm else ""

        # Data de vencimento + valor. Este é o par explicitamente associado no cabeçalho.
        vencimento = ""
        valor_cents = None
        for i, x in enumerate(lines):
            if "DATA VENCIMENTO" in x.upper() and i + 1 < len(lines):
                block = " ".join(lines[i:i+4])
                mv = re.search(r"(\d{2}/\d{2}/\d{4})\s+([0-9][0-9.,]*)", block)
                if mv:
                    vencimento = mv.group(1)
                    valor_cents = cents_from_text(mv.group(2))
                    break

        # Fallbacks de valor: TOTAL A PAGAR / TOTAL da conta.
        if valor_cents is None:
            patterns = [
                r"TOTAL\s+A\s+PAGAR\s*\(EM\s*R\$\)[\s\S]{0,100}?([0-9][0-9.,]*)",
                r"\bTOTAL\s+([0-9]{1,3}(?:\.[0-9]{3})*,\d{2})\b",
            ]
            for pat in patterns:
                m = re.search(pat, text, re.I)
                if m:
                    valor_cents = cents_from_text(m.group(1))
                    break

        if valor_cents is None:
            raise ValueError(f"Não foi possível localizar o valor da fatura na página {page_no}.")
        if not matricula:
            raise ValueError(f"Não foi possível localizar a matrícula na página {page_no}.")

        # Emissão: primeira data após o rótulo Data Emissão.
        emissao = ""
        for i, x in enumerate(lines):
            if x.upper() == "DATA EMISSÃO" and i + 1 < len(lines):
                for y in lines[i+1:i+10]:
                    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", y):
                        emissao = y
                        break
                break
        if not emissao:
            m = re.search(r"DATA\s+EMISS[ÃA]O[\s\S]{0,250}?(\d{2}/\d{2}/\d{4})", text, re.I)
            emissao = m.group(1) if m else ""

        records.append({
            "Página": page_no,
            "Matrícula": str(matricula),
            "Mês/Ano": mes_ano,
            "Data Emissão": emissao,
            "Data Vencimento": vencimento,
            "Valor_Centavos": int(valor_cents),
            "Nome na Fatura": extract_account_name(lines),
        })

    if not records:
        raise ValueError("Nenhuma fatura foi reconhecida no PDF.")

    df = pd.DataFrame(records)
    # Segurança institucional: não somar silenciosamente a mesma matrícula duas vezes
    # sem explicitar. Mantemos cada fatura em 'Faturas' e consolidamos por matrícula depois.
    df["Valor_Centavos"] = df["Valor_Centavos"].astype("int64")
    return df


# ---------------------------------------------------------------------------
# Histórico automático local
# ---------------------------------------------------------------------------
# O sistema mantém uma cópia acumulada das competências já processadas.
# Isso evita que um mês processado (ex.: Jul/2026) desapareça ao abrir o PDF
# do mês seguinte (ex.: Ago/2026), o que antes fazia o gráfico cair para zero.

def _history_month_columns(df: pd.DataFrame) -> List[str]:
    """Retorna somente colunas mensais válidas, ordenadas cronologicamente."""
    cols = [c for c in df.columns if re.fullmatch(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}", str(c))]
    return sorted(set(cols), key=month_sort_key)


def history_snapshot_from_control(control: pd.DataFrame) -> pd.DataFrame:
    """Extrai do controle apenas estrutura cadastral + histórico mensal.

    Colunas calculadas do mês atual (Aum / Red, % Centro e % Total) não são
    persistidas porque são sempre recalculadas a partir dos valores mensais.
    """
    if control is None or control.empty:
        return embedded_history_df()
    month_cols = _history_month_columns(control)
    keep = [c for c in ["Nº G", "LOCAL", "CEN", "MAT"] if c in control.columns] + month_cols
    hist = control[keep].copy()
    hist["MAT"] = hist["MAT"].astype(str)
    for c in month_cols:
        hist[c] = pd.to_numeric(hist[c], errors="coerce").astype("Int64")
    return hist


def _portable_history_path() -> Path:
    """Caminho principal do histórico: junto do script/executável, mantendo o pacote portátil."""
    if getattr(sys, "frozen", False):
        base = Path(sys.executable).resolve().parent
    else:
        base = Path(__file__).resolve().parent
    return base / "historico_conta_facil.json"


def _fallback_history_path() -> Path:
    """Caminho alternativo caso a pasta do programa não permita gravação."""
    root = Path(os.environ.get("APPDATA", str(Path.home())))
    return root / "Conta_Facil" / "historico_conta_facil.json"


def _serialize_history(df: pd.DataFrame) -> dict:
    month_cols = _history_month_columns(df)
    rows = []
    for _, row in df.iterrows():
        item = {}
        for c in ["Nº G", "LOCAL", "CEN", "MAT"] + month_cols:
            v = row.get(c, pd.NA)
            if pd.isna(v):
                item[c] = None
            elif c == "Nº G" or c in month_cols:
                item[c] = int(v)
            else:
                item[c] = str(v)
        rows.append(item)
    return {"version": 1, "rows": rows}


def save_local_history(df: pd.DataFrame) -> Path:
    """Salva o histórico acumulado sem interferir no processamento principal."""
    payload = _serialize_history(df)
    last_error = None
    for path in (_portable_history_path(), _fallback_history_path()):
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            tmp = path.with_suffix(path.suffix + ".tmp")
            tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            tmp.replace(path)
            return path
        except Exception as exc:
            last_error = exc
    raise OSError(f"Não foi possível salvar o histórico automático: {last_error}")


def load_local_history() -> Tuple[Optional[pd.DataFrame], Optional[Path]]:
    """Carrega o histórico automático salvo em execução anterior, se existir."""
    for path in (_portable_history_path(), _fallback_history_path()):
        if not path.exists():
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
            rows = payload.get("rows", [])
            if not rows:
                continue
            df = pd.DataFrame(rows)
            if not {"LOCAL", "CEN", "MAT"}.issubset(df.columns):
                continue
            if "Nº G" not in df.columns:
                df.insert(0, "Nº G", range(1, len(df) + 1))
            df["MAT"] = df["MAT"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
            for c in _history_month_columns(df):
                df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
            return df, path
        except Exception:
            # Um arquivo local corrompido não pode impedir a abertura do sistema.
            continue
    return None, None


def _center_month_total_or_nan(grp: pd.DataFrame, col: str) -> float:
    """Soma um centro no mês; retorna NaN quando o mês inteiro está ausente.

    NaN faz o Matplotlib deixar uma lacuna no gráfico. Isso é proposital:
    dado ausente não é faturamento zero.
    """
    vals = pd.to_numeric(grp[col], errors="coerce")
    total = vals.sum(min_count=1)
    return float("nan") if pd.isna(total) else float(total) / 100.0


# ---------------------------------------------------------------------------
# Transforma a BASE_HISTORICA embutida no código no mesmo formato tabular usado pelo restante da aplicação.
# ---------------------------------------------------------------------------
def embedded_history_df() -> pd.DataFrame:
    df = pd.DataFrame(BASE_HISTORICA)
    rename = {"ordem":"Nº G", "matricula":"MAT", "local":"LOCAL", "centro":"CEN"}
    df = df.rename(columns=rename)
    for c in [x for x in df.columns if "/" in x]:
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    df["MAT"] = df["MAT"].astype(str)
    return df


# ---------------------------------------------------------------------------
# Permite carregar um histórico externo em Excel e normaliza nomes de colunas, matrículas e valores para o modelo interno.
# ---------------------------------------------------------------------------
def normalize_history_excel(path: str) -> pd.DataFrame:
    """Lê um histórico Excel e normaliza para o modelo interno.

    Aceita tanto planilhas simples (cabeçalho na primeira linha) quanto os arquivos
    exportados pelo próprio Conta Fácil, que possuem um título nas primeiras linhas
    e o cabeçalho real na linha 3 do Excel.
    """
    xls = pd.ExcelFile(path)
    preferred = next((s for s in xls.sheet_names if s.lower() in {"controle", "planilha de controle"}), xls.sheet_names[0])

    # O Excel gerado pelo próprio sistema grava um título na linha 1 e deixa a
    # linha 2 em branco; nesse caso, os nomes das colunas estão na linha 3.
    # Para também aceitar planilhas externas comuns, tentamos localizar
    # automaticamente a linha que contém LOCAL/CEN/MAT.
    preview = pd.read_excel(path, sheet_name=preferred, header=None, nrows=12)
    header_row = None
    for idx, row in preview.iterrows():
        normalized = {str(v).strip().lower() for v in row.tolist() if not pd.isna(v)}
        has_local = any(v in {"local", "localização", "localizacao"} for v in normalized)
        has_center = any(v in {"cen", "centro"} for v in normalized)
        has_mat = any(v in {"mat", "matrícula", "matricula"} for v in normalized)
        if has_local and has_center and has_mat:
            header_row = int(idx)
            break

    # Fallback para planilhas tradicionais cuja primeira linha já é o cabeçalho.
    if header_row is None:
        header_row = 0

    raw = pd.read_excel(path, sheet_name=preferred, header=header_row)
    # Remove colunas completamente vazias, comuns quando há células mescladas/títulos.
    raw = raw.dropna(axis=1, how="all")

    colmap = {}
    for c in raw.columns:
        k = str(c).strip().lower()
        if k in {"nº g", "n° g", "no g", "ordem"}: colmap[c] = "Nº G"
        elif k in {"local", "localização", "localizacao"}: colmap[c] = "LOCAL"
        elif k in {"cen", "centro"}: colmap[c] = "CEN"
        elif k in {"mat", "matrícula", "matricula"}: colmap[c] = "MAT"
    raw = raw.rename(columns=colmap)

    required = {"LOCAL", "CEN", "MAT"}
    if not required.issubset(raw.columns):
        encontrados = ", ".join(map(str, raw.columns[:12]))
        raise ValueError(
            "A planilha histórica não foi reconhecida. É necessário existir uma aba "
            "Controle (ou equivalente) com as colunas LOCAL, CEN/Centro e MAT/Matrícula. "
            f"Colunas encontradas: {encontrados}"
        )

    if "Nº G" not in raw.columns:
        raw.insert(0, "Nº G", range(1, len(raw)+1))

    # Descarta eventuais linhas de rodapé/vazias sem matrícula.
    raw = raw[raw["MAT"].notna()].copy()
    raw["MAT"] = raw["MAT"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()

    # Identifica colunas mensais e converte os valores em reais do Excel para
    # centavos inteiros, formato usado internamente pelo sistema.
    month_cols = []
    canonical_names = {}
    month_pattern = re.compile(r"^(Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/(20\d{2})$", re.I)
    short_pattern = re.compile(r"^(Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)$", re.I)
    month_lookup = {m.lower(): m for m in MESES_PT.values()}
    for c in raw.columns:
        text = str(c).strip()
        m = month_pattern.fullmatch(text)
        if m:
            canon = f"{month_lookup[m.group(1).lower()]}/{m.group(2)}"
            month_cols.append(c)
            canonical_names[c] = canon
        elif short_pattern.fullmatch(text):
            month_cols.append(c)

    for c in month_cols:
        vals = pd.to_numeric(raw[c], errors="coerce")
        raw[c] = vals.apply(
            lambda x: pd.NA if pd.isna(x) else int((Decimal(str(x))*100).quantize(Decimal("1")))
        ).astype("Int64")

    if canonical_names:
        raw = raw.rename(columns=canonical_names)
        month_cols = [canonical_names.get(c, c) for c in month_cols]

    keep = ["Nº G", "LOCAL", "CEN", "MAT"] + month_cols
    return raw[keep].copy()


# ---------------------------------------------------------------------------
# Combina faturas do mês com o histórico por matrícula e calcula colunas de controle, variações e resumo por centro.
# ---------------------------------------------------------------------------
def build_control(faturas: pd.DataFrame, history: Optional[pd.DataFrame] = None) -> Tuple[pd.DataFrame, pd.DataFrame, str]:
    if history is None:
        history = embedded_history_df()
    history = history.copy()
    history["MAT"] = history["MAT"].astype(str)

    mes_series = faturas["Mês/Ano"].dropna().astype(str)
    mes_series = mes_series[mes_series.str.match(r"^(0[1-9]|1[0-2])/20\d{2}$")]
    if mes_series.empty:
        raise ValueError("Não foi possível identificar o mês/ano de referência nas faturas.")
    mes_ano = mes_series.mode().iloc[0]
    if (faturas["Mês/Ano"].fillna("") != mes_ano).any():
        distintos = sorted(set(faturas["Mês/Ano"].dropna().astype(str)))
        raise ValueError(f"O PDF contém meses de referência diferentes: {distintos}")
    current_col = month_key(mes_ano)

    atual = faturas.groupby("Matrícula", as_index=False, sort=False)["Valor_Centavos"].sum()
    atual = atual.rename(columns={"Matrícula":"MAT", "Valor_Centavos":current_col})
    atual["MAT"] = atual["MAT"].astype(str)

    # Adiciona matrículas novas sem descartar a estrutura histórica.
    base = history.merge(atual, on="MAT", how="outer", suffixes=("", "__novo"))
    if f"{current_col}__novo" in base.columns:
        if current_col in history.columns:
            base[current_col] = base[f"{current_col}__novo"].combine_first(base[current_col])
        else:
            base[current_col] = base[f"{current_col}__novo"]
        base = base.drop(columns=[f"{current_col}__novo"])

    # Para matrículas recém-descobertas, tenta usar o nome da própria fatura.
    names = faturas.drop_duplicates("Matrícula").set_index("Matrícula")["Nome na Fatura"].to_dict()
    max_order = pd.to_numeric(base.get("Nº G"), errors="coerce").max()
    max_order = int(max_order) if pd.notna(max_order) else 0
    missing_order = base["Nº G"].isna() if "Nº G" in base.columns else pd.Series(True, index=base.index)
    for idx in base.index[missing_order]:
        max_order += 1
        base.at[idx, "Nº G"] = max_order
    if "LOCAL" not in base: base["LOCAL"] = ""
    if "CEN" not in base: base["CEN"] = "NÃO MAPEADO"
    for idx, row in base.iterrows():
        if pd.isna(row.get("LOCAL")) or not str(row.get("LOCAL", "")).strip():
            base.at[idx, "LOCAL"] = names.get(str(row["MAT"]), "MATRÍCULA NOVA - REVISAR LOCAL") or "MATRÍCULA NOVA - REVISAR LOCAL"
        if pd.isna(row.get("CEN")) or not str(row.get("CEN", "")).strip():
            base.at[idx, "CEN"] = "NÃO MAPEADO"

    month_cols = [c for c in base.columns if re.fullmatch(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}", str(c))]
    month_cols = sorted(set(month_cols), key=month_sort_key)
    for c in month_cols:
        base[c] = pd.to_numeric(base[c], errors="coerce").astype("Int64")
    if current_col not in base.columns:
        base[current_col] = pd.Series([pd.NA]*len(base), dtype="Int64")
        month_cols.append(current_col)
        month_cols = sorted(set(month_cols), key=month_sort_key)

    # A variação institucional deve comparar com o mês calendário imediatamente anterior,
    # nunca com um mês antigo apenas porque houve lacuna no histórico.
    previous_col = previous_month_label(mes_ano)
    if previous_col not in base.columns:
        base[previous_col] = pd.Series([pd.NA]*len(base), dtype="Int64")
        month_cols.append(previous_col)
        month_cols = sorted(set(month_cols), key=month_sort_key)

    current_vals = base[current_col].fillna(0).astype("int64")
    total = int(current_vals.sum())
    center_totals = base.assign(_v=current_vals).groupby("CEN", dropna=False)["_v"].transform("sum")

    previous_has_data = base[previous_col].notna().any()
    if previous_has_data:
        prev_vals = base[previous_col].fillna(0).astype("int64")
        variation = []
        for cur, prev in zip(current_vals, prev_vals):
            if prev == 0:
                variation.append(None if cur == 0 else 100.0)
            else:
                variation.append(float((Decimal(int(cur-prev)) / Decimal(int(prev))) * Decimal(100)))
        base["Aum / Red"] = variation
    else:
        base["Aum / Red"] = [None]*len(base)

    base["% Centro"] = [float(Decimal(int(v))*100/Decimal(int(ct))) if ct else 0.0 for v,ct in zip(current_vals, center_totals)]
    base["% Total"] = [float(Decimal(int(v))*100/Decimal(total)) if total else 0.0 for v in current_vals]

    base["Nº G"] = pd.to_numeric(base["Nº G"], errors="coerce").fillna(0).astype(int)
    base = base.sort_values(["Nº G", "MAT"], kind="stable").reset_index(drop=True)
    cols = ["Nº G", "LOCAL", "CEN", "MAT"] + month_cols + ["Aum / Red", "% Centro", "% Total"]
    base = base[cols]

    # Resumo por centro calculado pelo pandas sobre centavos inteiros.
    work = base[["CEN", current_col]].copy()
    work[current_col] = work[current_col].fillna(0).astype("int64")
    resumo = work.groupby("CEN", as_index=False)[current_col].sum().rename(columns={current_col:"Valor_Centavos"})
    if previous_has_data:
        prev = base[["CEN", previous_col]].copy()
        prev[previous_col] = prev[previous_col].fillna(0).astype("int64")
        prev = prev.groupby("CEN", as_index=False)[previous_col].sum().rename(columns={previous_col:"Anterior_Centavos"})
        resumo = resumo.merge(prev, on="CEN", how="left")
        resumo["Variação %"] = resumo.apply(
            lambda r: None if int(r["Anterior_Centavos"]) == 0 and int(r["Valor_Centavos"]) == 0
            else (100.0 if int(r["Anterior_Centavos"]) == 0 else float(Decimal(int(r["Valor_Centavos"])-int(r["Anterior_Centavos"])) * 100 / Decimal(int(r["Anterior_Centavos"])))), axis=1)
    else:
        resumo["Anterior_Centavos"] = 0
        resumo["Variação %"] = None
    resumo["% Total"] = resumo["Valor_Centavos"].apply(lambda v: float(Decimal(int(v))*100/Decimal(total)) if total else 0.0)
    order = {c:i for i,c in enumerate(CENTROS_ORDEM)}
    resumo["_ord"] = resumo["CEN"].map(order).fillna(999)
    resumo = resumo.sort_values(["_ord", "CEN"]).drop(columns="_ord").reset_index(drop=True)

    return base, resumo, mes_ano


# ---------------------------------------------------------------------------
# Exporta os três conjuntos principais (controle, faturas e resumo) para um arquivo Excel formatado.
# ---------------------------------------------------------------------------
def export_excel(path: str, control: pd.DataFrame, faturas: pd.DataFrame, resumo: pd.DataFrame, mes_ano: str):
    current_col = month_key(mes_ano)
    # Exportamos valores em reais para leitura humana no Excel; cálculos já ocorreram
    # internamente em centavos exatos.
    ctrl = control.copy()
    month_cols = [c for c in ctrl.columns if re.fullmatch(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}", str(c))]
    for c in month_cols:
        ctrl[c] = ctrl[c].apply(lambda x: None if pd.isna(x) else int(x)/100.0)
    fat = faturas.copy()
    fat["Valor (R$)"] = fat["Valor_Centavos"].astype("int64") / 100.0
    fat = fat.drop(columns=["Valor_Centavos"])
    res = resumo.copy()
    res["Valor Atual (R$)"] = res["Valor_Centavos"].astype("int64") / 100.0
    res["Valor Anterior (R$)"] = res["Anterior_Centavos"].astype("int64") / 100.0
    res = res.drop(columns=["Valor_Centavos", "Anterior_Centavos"])

    with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
        ctrl.to_excel(writer, sheet_name="Controle", index=False, startrow=2)
        fat.to_excel(writer, sheet_name="Faturas", index=False, startrow=2)
        res.to_excel(writer, sheet_name="Resumo Centros", index=False, startrow=2)

        wb = writer.book
        title_fmt = wb.add_format({"bold":True, "font_size":14, "align":"center", "valign":"vcenter", "bg_color":"#1F4E78", "font_color":"white"})
        header_fmt = wb.add_format({"bold":True, "bg_color":"#D9EAF7", "border":1, "align":"center", "valign":"vcenter", "text_wrap":True})
        money_fmt = wb.add_format({"num_format":'R$ #,##0.00', "border":1})
        pct_fmt = wb.add_format({"num_format":'0.0%', "border":1})
        pct100_fmt = wb.add_format({"num_format":'0.0', "border":1})
        text_fmt = wb.add_format({"border":1, "valign":"top"})
        warn_fmt = wb.add_format({"bg_color":"#FFF2CC", "font_color":"#9C6500"})

        for sheet_name, df in [("Controle", ctrl), ("Faturas", fat), ("Resumo Centros", res)]:
            ws = writer.sheets[sheet_name]
            ws.merge_range(0, 0, 0, max(0, len(df.columns)-1), f"UFRB - EMBASA | Referência {mes_ano}", title_fmt)
            ws.set_row(0, 24)
            for col_idx, name in enumerate(df.columns):
                ws.write(2, col_idx, name, header_fmt)
            ws.freeze_panes(3, 0)
            ws.autofilter(2, 0, 2+len(df), max(0, len(df.columns)-1))

        ws = writer.sheets["Controle"]
        for i,c in enumerate(ctrl.columns):
            width = 12
            if c == "LOCAL": width = 43
            elif c in {"CEN", "MAT"}: width = 14
            elif c == "Nº G": width = 7
            ws.set_column(i, i, width, text_fmt)
            if c in month_cols:
                ws.set_column(i, i, 14, money_fmt)
            elif c in {"Aum / Red", "% Centro", "% Total"}:
                # Valores estão no DataFrame em percentuais 0..100, então formato simples.
                ws.set_column(i, i, 12, pct100_fmt)
        cen_col = ctrl.columns.get_loc("CEN")
        ws.conditional_format(3, cen_col, 2+len(ctrl), cen_col, {"type":"text", "criteria":"containing", "value":"NÃO MAPEADO", "format":warn_fmt})

        ws = writer.sheets["Faturas"]
        for i,c in enumerate(fat.columns):
            width = 16
            if c == "Nome na Fatura": width = 45
            ws.set_column(i, i, width, money_fmt if c == "Valor (R$)" else text_fmt)

        ws = writer.sheets["Resumo Centros"]
        for i,c in enumerate(res.columns):
            width = 20
            fmt = money_fmt if "(R$)" in c else text_fmt
            if c in {"Variação %", "% Total"}: fmt = pct100_fmt
            ws.set_column(i, i, width, fmt)

        # Aba de memória de cálculo clara e auditável.
        mem = wb.add_worksheet("Memória de Cálculo")
        mem.write("A1", "MEMÓRIA DE CÁLCULO", title_fmt)
        mem.write("A3", "Regra")
        mem.write("B3", "Descrição")
        rows = [
            ("Valor total", "Soma dos valores das faturas em centavos inteiros; apresentação ÷ 100."),
            ("Aum / Red", "((mês atual - mês anterior) / mês anterior) × 100."),
            ("% Centro", "valor da matrícula / soma do centro × 100."),
            ("% Total", "valor da matrícula / total geral UFRB × 100."),
        ]
        for r,(a,b) in enumerate(rows, start=3):
            mem.write(r,0,a); mem.write(r,1,b)
        mem.set_column(0,0,20); mem.set_column(1,1,75)


# ---------------------------------------------------------------------------
# Aplica cor de fundo em células do Word usando o XML subjacente do python-docx.
# ---------------------------------------------------------------------------
def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


# ---------------------------------------------------------------------------
# Padroniza texto, tamanho e negrito em células de tabelas do Word.
# ---------------------------------------------------------------------------
def set_cell_text(cell, text: str, bold=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


# ---------------------------------------------------------------------------
# Cria uma tabela genérica no Word para seções resumidas do relatório.
# ---------------------------------------------------------------------------
def add_doc_table(doc: Document, headers: List[str], rows: List[List[str]], widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i,h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=font_size)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i,v in enumerate(row):
            set_cell_text(cells[i], v, size=font_size)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return table


# ---------------------------------------------------------------------------
# Ajusta margens internas das células da planilha detalhada para aproveitar melhor a página em paisagem.
# ---------------------------------------------------------------------------
def _set_docx_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    """Margens internas compactas para a planilha detalhada do relatório."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


# ---------------------------------------------------------------------------
# Marca a primeira linha da tabela para repetir o cabeçalho em cada nova página do Word.
# ---------------------------------------------------------------------------
def _set_docx_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


# ---------------------------------------------------------------------------
# Evita que uma linha da tabela seja dividida entre duas páginas do relatório.
# ---------------------------------------------------------------------------
def _set_docx_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


# ---------------------------------------------------------------------------
# Aplica a tipografia/alinhamento padronizados da planilha de controle no Word.
# ---------------------------------------------------------------------------
def _format_control_cell(cell, text, *, bold=False, size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Formatação visual da planilha de controle no Word; não altera valores."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    # Garante Arial também em renderizadores que consultam fontes East Asia.
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_docx_cell_margins(cell)


# ---------------------------------------------------------------------------
# Monta a planilha detalhada do ateste com larguras fixas e formatação própria para Word/Google Docs.
# ---------------------------------------------------------------------------
def add_control_doc_table(doc: Document, headers: List[str], rows: List[List[str]]):
    """Tabela detalhada do ateste, otimizada para leitura em página paisagem.

    Apenas apresentação: fonte, alinhamento, largura, margens e repetição de cabeçalho.
    """
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    # Layout fixo impede o Word/Google Docs de espremer colunas e quebrar cabeçalhos.
    tblPr = table._tbl.tblPr
    tblLayout = tblPr.first_child_found_in("w:tblLayout")
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")

    # Larguras pensadas para A4 paisagem com margens de 1,2 cm.
    month_re = re.compile(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}")
    widths = []
    for h in headers:
        if h == "Nº G": widths.append(0.9)
        elif h == "LOCAL": widths.append(6.6)
        elif h == "CEN": widths.append(1.2)
        elif h == "MAT": widths.append(1.9)
        elif month_re.fullmatch(str(h)): widths.append(1.85)
        elif h == "Aum / Red": widths.append(1.55)
        elif h in {"% Centro", "% Total"}: widths.append(1.40)
        else: widths.append(1.45)

    header_row = table.rows[0]
    _set_docx_repeat_header(header_row)
    _set_docx_cant_split(header_row)
    for i, h in enumerate(headers):
        _format_control_cell(header_row.cells[i], h, bold=True, size=8.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(header_row.cells[i], "D9EAF7")
        header_row.cells[i].width = Cm(widths[i])

    for values in rows:
        row = table.add_row()
        _set_docx_cant_split(row)
        for i, value in enumerate(values):
            h = headers[i]
            if h == "LOCAL":
                align = WD_ALIGN_PARAGRAPH.LEFT
            elif month_re.fullmatch(str(h)) or h in {"Aum / Red", "% Centro", "% Total"}:
                align = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                align = WD_ALIGN_PARAGRAPH.CENTER
            _format_control_cell(row.cells[i], value, size=8.0, align=align)
            row.cells[i].width = Cm(widths[i])

    return table


# ---------------------------------------------------------------------------
# Gera imagens temporárias dos gráficos usados posteriormente dentro do relatório Word.
# ---------------------------------------------------------------------------
def make_chart_images(control: pd.DataFrame, resumo: pd.DataFrame, mes_ano: str, out_dir: str) -> Tuple[str, str]:
    current_col = month_key(mes_ano)
    os.makedirs(out_dir, exist_ok=True)
    center_path = os.path.join(out_dir, "grafico_centros.png")
    trend_path = os.path.join(out_dir, "grafico_tendencia.png")

    fig = Figure(figsize=(7.2, 3.8), tight_layout=True)
    ax = fig.add_subplot(111)
    data = resumo.copy()
    ax.bar(data["CEN"], data["Valor_Centavos"] / 100.0)
    ax.set_title(f"Distribuição do faturamento por centro - {mes_ano}")
    ax.set_ylabel("Valor (R$)")
    ax.tick_params(axis="x", rotation=30)
    fig.savefig(center_path, dpi=180)

    month_cols = [c for c in control.columns if re.fullmatch(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}", str(c))]
    month_cols = sorted(month_cols, key=month_sort_key)
    fig2 = Figure(figsize=(7.2, 3.8), tight_layout=True)
    ax2 = fig2.add_subplot(111)
    for cen, grp in control.groupby("CEN"):
        if cen == "NÃO MAPEADO":
            continue
        ys = [_center_month_total_or_nan(grp, c) for c in month_cols]
        ax2.plot(month_cols, ys, marker="o", label=cen)
    ax2.set_title("Evolução mensal do faturamento por centro")
    ax2.set_ylabel("Valor (R$)")
    ax2.tick_params(axis="x", rotation=30)
    ax2.legend(fontsize=7, ncol=3)
    fig2.savefig(trend_path, dpi=180)
    return center_path, trend_path




# ---------------------------------------------------------------------------
# Converte valores opcionais em float de forma defensiva para a geração de textos analíticos.
# ---------------------------------------------------------------------------
def _safe_float(value):
    try:
        if value is None or pd.isna(value):
            return None
        return float(value)
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Formata listas de centros e percentuais em texto corrido para o relatório.
# ---------------------------------------------------------------------------
def _fmt_center_list(items: List[Tuple[str, float]]) -> str:
    if not items:
        return "nenhum centro"
    parts = [f"{c} ({pct(v)})" for c, v in items]
    if len(parts) == 1:
        return parts[0]
    return ", ".join(parts[:-1]) + " e " + parts[-1]


# ---------------------------------------------------------------------------
# Gera os parágrafos analíticos com base exclusivamente nos números calculados e nas ocorrências informadas pelo usuário.
# ---------------------------------------------------------------------------
def build_analytical_narrative(control: pd.DataFrame, resumo: pd.DataFrame, mes_ano: str,
                               ocorrencias: str = "") -> Dict[str, object]:
    """Gera texto analítico a partir dos dados, sem inventar causas operacionais."""
    current_col = month_key(mes_ano)
    prev_col = previous_month_label(mes_ano)
    total = int(control[current_col].fillna(0).astype('int64').sum()) if current_col in control.columns else 0
    prev_has_data = prev_col in control.columns and control[prev_col].notna().any()
    prev_total = int(control[prev_col].fillna(0).astype('int64').sum()) if prev_has_data else 0
    geral_var = ((total - prev_total) * 100.0 / prev_total) if prev_has_data and prev_total else None

    center_rows = []
    for _, r in resumo.iterrows():
        cen = str(r.get('CEN', ''))
        var = _safe_float(r.get('Variação %'))
        share = _safe_float(r.get('% Total')) or 0.0
        value = int(r.get('Valor_Centavos', 0) or 0)
        center_rows.append((cen, var, share, value))

    increases = sorted([(c, v) for c, v, _, _ in center_rows if v is not None and v > 0.5], key=lambda x: x[1], reverse=True)
    reductions = sorted([(c, v) for c, v, _, _ in center_rows if v is not None and v < -0.5], key=lambda x: x[1])
    stable = sorted([(c, v) for c, v, _, _ in center_rows if v is not None and -0.5 <= v <= 0.5], key=lambda x: x[0])

    reference = month_long(mes_ano)
    if geral_var is None:
        panorama = (
            f"No mês de {reference}, foram consolidados os valores das contas de água da UFRB. "
            f"Não foi possível calcular a variação em relação ao mês imediatamente anterior ({prev_col}), "
            "pois a base histórica carregada não contém dados suficientes para essa competência. "
            f"O faturamento total apurado no período foi de {brl(total)}."
        )
    else:
        movimento = "aumento" if geral_var > 0.5 else "redução" if geral_var < -0.5 else "estabilidade"
        panorama = (
            f"No mês de {reference}, observou-se {movimento} de {pct(abs(geral_var))} no faturamento de água da UFRB "
            f"em relação ao mês anterior, passando de {brl(prev_total)} para {brl(total)}. "
        )
        if increases:
            panorama += f"Apresentaram aumento: {_fmt_center_list(increases)}. "
        if reductions:
            panorama += f"Apresentaram redução: {_fmt_center_list(reductions)}. "
        if stable:
            panorama += f"Permaneceram relativamente estáveis: {_fmt_center_list(stable)}. "
        panorama += "A participação de cada centro no total institucional é apresentada na tabela e nos gráficos a seguir."

    center_texts = []
    for cen, var, share, value in sorted(center_rows, key=lambda x: x[3], reverse=True):
        grp = control[control['CEN'].astype(str) == cen].copy()
        if grp.empty:
            continue
        txt = f"O {cen} totalizou {brl(value)} no período, correspondendo a {pct(share)} do faturamento total da UFRB."
        if var is not None:
            if var > 0.5:
                txt += f" Em comparação ao mês anterior, o centro apresentou aumento de {pct(var)}."
            elif var < -0.5:
                txt += f" Em comparação ao mês anterior, o centro apresentou redução de {pct(abs(var))}."
            else:
                txt += " Em comparação ao mês anterior, o centro apresentou relativa estabilidade."

        if prev_has_data and current_col in grp.columns and prev_col in grp.columns:
            tmp = grp[['MAT', 'LOCAL', current_col, prev_col]].copy()
            tmp = tmp[tmp[prev_col].notna() & tmp[current_col].notna()]
            tmp = tmp[tmp[prev_col].astype('int64') > 0]
            if not tmp.empty:
                tmp['var'] = (tmp[current_col].astype('int64') - tmp[prev_col].astype('int64')) * 100.0 / tmp[prev_col].astype('int64')
                tmp['impacto'] = (tmp[current_col].astype('int64') - tmp[prev_col].astype('int64')).abs()
                destaques = tmp.sort_values('impacto', ascending=False).head(2)
                desc=[]
                for _, d in destaques.iterrows():
                    direcao = 'aumento' if d['var'] > 0 else 'redução'
                    desc.append(f"matrícula {d['MAT']} ({str(d['LOCAL']).strip()}), com {direcao} de {pct(abs(float(d['var'])))}")
                if desc:
                    txt += " Entre as maiores variações financeiras do centro, destacaram-se " + " e ".join(desc) + "."
        center_texts.append((cen, txt))

    if geral_var is None:
        conclusion = (
            f"Conclui-se que o faturamento apurado para {reference} foi de {brl(total)}. "
            "Como não há dados do mês imediatamente anterior na base carregada, recomenda-se completar o histórico antes de concluir sobre tendência de aumento ou redução."
        )
    else:
        mov = "aumento" if geral_var > 0.5 else "redução" if geral_var < -0.5 else "estabilidade"
        conclusion = (
            f"Conclui-se que houve {mov} geral de {pct(abs(geral_var))} no faturamento de água da UFRB no mês de {reference}. "
            "Recomenda-se manter o acompanhamento periódico das matrículas com maior impacto financeiro e das variações relevantes por centro."
        )
    if ocorrencias.strip():
        conclusion += " Conforme registro informado pelo gestor para o período: " + ocorrencias.strip()

    return {
        'geral_var': geral_var,
        'panorama': panorama,
        'centros': center_texts,
        'conclusao': conclusion,
        'grande_variacao': geral_var is not None and abs(geral_var) >= 20,
    }


# ---------------------------------------------------------------------------
# Orquestra a criação completa do Relatório Mensal de Ateste em Word, incluindo texto, tabelas, gráficos e assinatura.
# ---------------------------------------------------------------------------
def generate_word_report(path: str, control: pd.DataFrame, faturas: pd.DataFrame, resumo: pd.DataFrame,
                         mes_ano: str, contrato: str, gestor: str, siape: str, data_ateste: str, ocorrencias: str = ""):
    doc = Document()
    sec = doc.sections[0]
    # Formatação acadêmico-institucional baseada nas convenções ABNT: A4,
    # margens 3 cm (superior/esquerda) e 2 cm (inferior/direita), fonte 12,
    # texto justificado, recuo de 1,25 cm e entrelinhas 1,5.
    # A lógica de cálculos, extração e análise permanece inalterada.
    sec.top_margin = Cm(3); sec.bottom_margin = Cm(2); sec.left_margin = Cm(3); sec.right_margin = Cm(2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    def aplicar_abnt_documento(document):
        """Normaliza parágrafos narrativos sem alterar conteúdo ou cálculos."""
        titulos = {
            "Introdução", "Panorama Geral do Consumo", "Análise por Centro",
            "Conclusão", "Verificação de conformidade", "DEMAIS OBSERVAÇÕES:"
        }
        for par in document.paragraphs:
            texto = par.text.strip()
            if not texto:
                continue
            # Cabeçalhos, títulos, legendas e assinatura conservam alinhamento próprio.
            if (texto in titulos or texto.startswith("Gráfico ") or
                texto.startswith("PLANILHA DE CONTROLE") or
                texto.startswith("UNIVERSIDADE FEDERAL") or
                texto.startswith("PRÓ-REITORIA") or
                texto.startswith("COORDENADORIA") or
                texto.startswith("RELATÓRIO MENSAL") or
                texto == "Contratos Natureza Continuada" or
                texto.startswith("DECLARO QUE") or
                texto == "NOME DO GESTOR DO CONTRATO - SIAPE"):
                par.paragraph_format.first_line_indent = Cm(0)
                continue
            # Campos administrativos e perguntas do formulário não recebem recuo.
            if (texto.startswith(("Contrato:", "Empresa:", "Mês de Referência:",
                                  "Faturamento Recebido em:", "1 -", "2 -", "2.1 -",
                                  "3 -", "3.1 -", "( X )", "(   )", "FATURA(S):",
                                  "Data:", "Total geral conferido:", "ATENÇÃO:")) or
                texto.startswith("_") or texto.startswith("As matrículas extraídas")):
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                continue
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.first_line_indent = Cm(1.25)
            par.paragraph_format.line_spacing = 1.5
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)

    for text, size in [("UNIVERSIDADE FEDERAL DO RECÔNCAVO DA BAHIA", 12), ("PRÓ-REITORIA DE ADMINISTRAÇÃO", 12), ("COORDENADORIA DE CONTRATOS", 11)]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(text); r.bold=True; r.font.size=Pt(size)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("RELATÓRIO MENSAL DE ATESTE"); r.bold=True; r.font.size=Pt(13)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Contratos Natureza Continuada"); r.bold=True

    doc.add_paragraph(f"Contrato: {contrato}")
    doc.add_paragraph("Empresa: EMBASA")
    doc.add_paragraph(f"Mês de Referência: {month_long(mes_ano).upper()}")
    doc.add_paragraph(f"Faturamento Recebido em: {data_ateste}")

    p=doc.add_paragraph(); r=p.add_run("1 - Consta nos documentos que estão sendo enviados no processo de pagamento:"); r.bold=True
    doc.add_paragraph("FATURAS, Relatório de Ateste e Planilha Memória de Cálculo")
    p=doc.add_paragraph(); r=p.add_run("2 - A Contratada realizou algum desconto/acréscimo indevido na fatura do presente mês?"); r.bold=True
    doc.add_paragraph("( X ) Não    (   ) Sim")
    doc.add_paragraph("2.1 - Indique qual acréscimo/desconto e as medidas adotadas: ______________________________________")

    total = int(faturas["Valor_Centavos"].sum())
    current_col = month_key(mes_ano)
    month_cols = sorted([c for c in control.columns if re.fullmatch(r"(?:Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/20\d{2}", str(c))], key=month_sort_key)
    prev_col = previous_month_label(mes_ano)
    previous_has_data = prev_col in control.columns and control[prev_col].notna().any()
    previous_total = int(control[prev_col].fillna(0).astype("int64").sum()) if previous_has_data else 0
    if previous_has_data and previous_total:
        geral_var = float(Decimal(total-previous_total)*100/Decimal(previous_total))
    else:
        geral_var = None

    p=doc.add_paragraph(); r=p.add_run("Introdução"); r.bold=True
    doc.add_paragraph(
        f"O presente relatório tem por objetivo analisar o faturamento das contas de água da Universidade Federal do Recôncavo da Bahia (UFRB), "
        f"referente ao mês de {month_long(mes_ano)}, com base nas faturas emitidas pela EMBASA. Os dados foram consolidados por matrícula e centro "
        f"para subsidiar a conferência, o ateste e o acompanhamento institucional."
    )
    analise = build_analytical_narrative(control, resumo, mes_ano, ocorrencias)
    p=doc.add_paragraph(); r=p.add_run("Panorama Geral do Consumo"); r.bold=True
    doc.add_paragraph(analise["panorama"])

    rows=[]
    for _,r in resumo.iterrows():
        rows.append([str(r["CEN"]), brl(r["Valor_Centavos"]), pct(r["Variação %"]), pct(r["% Total"])])
    add_doc_table(doc, ["Centro", "Valor Atual", "Aum/Red", "% Total"], rows, font_size=9)

    temp_dir = str(Path(path).with_suffix("")) + "_tmp"
    center_img, trend_img = make_chart_images(control, resumo, mes_ano, temp_dir)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); r.add_picture(center_img, width=Cm(15.5))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=p.add_run("Gráfico 1. Distribuição do faturamento por centro."); rr.italic=True; rr.font.size=Pt(9)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); r.add_picture(trend_img, width=Cm(15.5))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=p.add_run("Gráfico 2. Evolução mensal do faturamento por centro."); rr.italic=True; rr.font.size=Pt(9)

    p=doc.add_paragraph(); r=p.add_run("Análise por Centro"); r.bold=True
    for cen, texto in analise["centros"]:
        p=doc.add_paragraph();
        rr=p.add_run(f"{cen}: "); rr.bold=True
        p.add_run(texto)

    p=doc.add_paragraph(); r=p.add_run("Conclusão"); r.bold=True
    doc.add_paragraph(analise["conclusao"])

    # Tabela detalhada em seção paisagem para aproximar a planilha oficial.
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.orientation = 1  # landscape
    new_sec.page_width, new_sec.page_height = sec.page_height, sec.page_width
    new_sec.top_margin = Cm(1.2); new_sec.bottom_margin = Cm(1.2); new_sec.left_margin = Cm(1.2); new_sec.right_margin = Cm(1.2)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r=p.add_run(f"PLANILHA DE CONTROLE - {month_long(mes_ano).upper()}"); r.bold=True
    r.font.name = "Arial"; r.font.size = Pt(12)
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    display_cols = ["Nº G", "LOCAL", "CEN", "MAT"] + month_cols[-5:] + ["Aum / Red", "% Centro", "% Total"]
    detail=[]
    for _,row in control.iterrows():
        vals=[]
        for c in display_cols:
            if c in month_cols:
                vals.append("" if pd.isna(row[c]) else brl(int(row[c])).replace("R$ ", ""))
            elif c in {"Aum / Red", "% Centro", "% Total"}:
                vals.append(pct(row[c]))
            else:
                vals.append(str(row[c]))
        detail.append(vals)
    # DESIGN ONLY: tabela detalhada em Arial 8, largura fixa e alinhamentos próprios.
    # Os dados e a lógica permanecem exatamente os mesmos.
    add_control_doc_table(doc, display_cols, detail)

    # Retorna a retrato para encerramento/ateste.
    end_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    end_sec.orientation = 0
    end_sec.page_width, end_sec.page_height = sec.page_width, sec.page_height
    end_sec.top_margin = Cm(1.5); end_sec.bottom_margin = Cm(1.5); end_sec.left_margin = Cm(1.8); end_sec.right_margin = Cm(1.8)

    p=doc.add_paragraph(); r=p.add_run("Verificação de conformidade"); r.bold=True
    unmapped = int((control["CEN"] == "NÃO MAPEADO").sum())
    if unmapped == 0:
        doc.add_paragraph("As matrículas extraídas do PDF foram cruzadas com a estrutura de controle conhecida, sem pendências de centro/localização.")
    else:
        doc.add_paragraph(f"ATENÇÃO: existem {unmapped} matrícula(s) sem mapeamento institucional de centro/localização. Recomenda-se regularizar a planilha antes do ateste definitivo.")
    doc.add_paragraph(f"Total geral conferido: {brl(total)}. Quantidade de faturas: {len(faturas)}.")

    p=doc.add_paragraph(); r=p.add_run("3 - Foi verificada alguma grande variação no faturamento do presente mês?"); r.bold=True
    grande = bool(analise["grande_variacao"])
    doc.add_paragraph("( X ) Sim    (   ) Não" if grande else "(   ) Sim    ( X ) Não")
    if ocorrencias.strip():
        doc.add_paragraph("3.1 - Indique o ocorrido: " + ocorrencias.strip())
    else:
        doc.add_paragraph("3.1 - Indique o ocorrido: _________________________________________________________________")
    doc.add_paragraph("DEMAIS OBSERVAÇÕES:")
    doc.add_paragraph("________________________________________________________________________________________")
    doc.add_paragraph("________________________________________________________________________________________")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("DECLARO QUE OS SERVIÇOS FORAM INTEGRALMENTE PRESTADOS"); r.bold=True
    doc.add_paragraph(f"FATURA(S): {len(faturas)} conta(s)")
    doc.add_paragraph(f"Data: {data_ateste}")
    doc.add_paragraph("\n\n____________________________________________")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f"{gestor.upper()} - {siape}"); r.bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("NOME DO GESTOR DO CONTRATO - SIAPE")

    # Aplica a normalização ABNT somente à apresentação textual.
    # Nenhuma regra de negócio, valor, variação ou classificação é modificada.
    aplicar_abnt_documento(doc)
    doc.save(path)
    # Limpeza dos PNGs temporários após o Word ter incorporado as imagens.
    try:
        import shutil
        shutil.rmtree(temp_dir, ignore_errors=True)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Adaptador entre pandas.DataFrame e QTableView. Ele controla como cada célula, cabeçalho, alinhamento e valor são apresentados na planilha da interface.
# ---------------------------------------------------------------------------